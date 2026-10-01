"""Gera a planilha de controle NOMAD puffs (Excel / Google Sheets).

Uso: python planilha/gerar_planilha.py  ->  planilha/NOMAD-puffs-controle.xlsx
Fonte dos dados: docs/catalogo.pdf (preços, sabores e taxas de entrega).
"""
import os
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SAIDA = Path(os.environ.get("NOMAD_SAIDA") or Path(__file__).with_name("NOMAD-puffs-controle.xlsx"))
LINHAS_VENDAS = int(os.environ.get("NOMAD_LINHAS", 1000))     # linhas de VENDAS já preparadas com fórmulas
LINHAS_ENTREGAS = int(os.environ.get("NOMAD_LINHAS", 500))    # linhas do registro de entregas
FIM_VENDAS = 4 + LINHAS_VENDAS


def vcol(col):
    """Intervalo limitado de uma coluna de VENDAS (mais leve que a coluna inteira)."""
    return f"VENDAS!${col}$5:${col}${FIM_VENDAS}"

MODELOS = [
    ("V55", 85, ["Pineapple Ice", "Uva Ice", "Icy Mint"]),
    ("V155", 110, ["Pineapple Ice", "Menthol", "Grape Ice", "Watermelon Ice", "Icy Mint"]),
    ("V400 Mix Slim", 140, ["Icy Mint + Peach Grape", "Menthol + Mighty Melon",
                            "Mango + Passion Fruit Guava", "Strawberry Grape Ice + Kiwi Watermelon",
                            "Cherry + Grape"]),
    ("Elfbar Pro 40K", 140, ["Sour Apple Ice", "Strawberry Blend", "Pink Lemonade",
                             "Watermelon + Peach Frost", "Tropical Baja"]),
]
REGIOES = [
    ("Quadras 700 Sul a 200 Norte/Sul", "Quadras", 8),
    ("Quadras 300 Norte a 600 Norte", "Quadras", 10),
    ("Quadras 800 a 1200", "Quadras", 10),
    ("Quadras 1300 a 1500 Sul", "Quadras", 15),
    ("Santo Amaro", "Bairro", 15),
    ("Lago Norte", "Bairro", 20),
    ("Bertaville e Aurenys", "Bairro", 30),
    ("Taquaralto e Lago Sul", "Bairro", 35),
    ("Taquari", "Bairro", 35),
    ("Retirada / sem entrega", "Outro", 0),
]
PAGAMENTOS = ["PIX", "Cartão de débito", "Cartão de crédito"]
STATUS = ["Entregue", "Em rota", "Pendente", "Cancelado"]
ENTREGADORES = ["Motoboy 1", "Motoboy 2", "Motoboy 3"]

# ---- identidade visual NOMAD -------------------------------------------------
PRETO, CREME, LARANJA, CINZA = "131416", "E8DCC4", "E38A4E", "2A2C30"
FONTE = "Arial"
f_titulo = Font(name=FONTE, size=16, bold=True, color=CREME)
f_sub = Font(name=FONTE, size=10, italic=True, color=LARANJA)
f_cab = Font(name=FONTE, size=10, bold=True, color=PRETO)
f_txt = Font(name=FONTE, size=10, color="000000")
f_input = Font(name=FONTE, size=10, color="0000FF")          # azul = digitado
f_link = Font(name=FONTE, size=10, color="008000")           # verde = vem de outra aba
f_total = Font(name=FONTE, size=10, bold=True, color="000000")
fill_titulo = PatternFill("solid", fgColor=PRETO)
fill_cab = PatternFill("solid", fgColor=LARANJA)
fill_input = PatternFill("solid", fgColor="FFF7D6")          # amarelo claro = preencher
fill_total = PatternFill("solid", fgColor=CREME)
fina = Side(style="thin", color="BFBFBF")
borda = Border(left=fina, right=fina, top=fina, bottom=fina)
MOEDA = '"R$" #,##0.00;-"R$" #,##0.00;"-"'
INTEIRO = '#,##0;-#,##0;"-"'
DATA = "dd/mm/yyyy"
centro = Alignment(horizontal="center", vertical="center", wrap_text=True)


def titulo(ws, texto, sub, ate_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ate_col)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ate_col)
    ws["A1"], ws["A2"] = texto, sub
    for c in range(1, ate_col + 1):
        ws.cell(1, c).fill = fill_titulo
        ws.cell(2, c).fill = fill_titulo
    ws["A1"].font, ws["A2"].font = f_titulo, f_sub
    ws.row_dimensions[1].height = 28
    ws.sheet_view.showGridLines = False


def cabecalho(ws, linha, col, nomes, larguras=None):
    for i, nome in enumerate(nomes):
        c = ws.cell(linha, col + i, nome)
        c.font, c.fill, c.alignment, c.border = f_cab, fill_cab, centro, borda
        if larguras:
            ws.column_dimensions[get_column_letter(col + i)].width = larguras[i]
    ws.row_dimensions[linha].height = 30


def cel(ws, ref, valor, fonte=f_txt, fmt=None, fill=None):
    c = ws[ref]
    c.value, c.font, c.border = valor, fonte, borda
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def lista(ws, formula, ref, prompt):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True,
                        showErrorMessage=True, errorTitle="Valor inválido",
                        error="Escolha uma opção da lista.", promptTitle="Escolha", prompt=prompt)
    ws.add_data_validation(dv)
    dv.add(ref)


wb = Workbook()

# ============================ INÍCIO (legenda) ================================
ws = wb.active
ws.title = "INÍCIO"
titulo(ws, "NOMAD puffs — Controle de vendas", "VAPOR SEM FRONTEIRAS · Palmas – TO", 4)
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 90
linhas = [
    ("COMO USAR", ""),
    ("1. VENDAS", "Lance cada item vendido em uma linha. Itens do mesmo pedido repetem o Nº do pedido, "
                  "a data, o pagamento, a região e o entregador. A taxa de entrega só é cobrada na 1ª linha do pedido."),
    ("2. ENTREGADOR", "Cadastre os motoboys. No registro, digite só o Nº do pedido: data, motoboy, região, "
                      "modelos e valores aparecem sozinhos. Atualize o status da entrega."),
    ("3. MERCADORIA", "Digite as entradas de estoque (compras) por sabor. Vendido e estoque atual são automáticos."),
    ("4. VALOR", "Preço de venda de cada modelo (vem do catálogo). Preencha o custo para ver a margem."),
    ("5. MAPA", "Taxas de entrega por quadra/bairro de Palmas e o desempenho de cada região."),
    ("6. RESUMO", "Painel com os totais: faturamento, pedidos, ticket médio, PIX x cartão, modelos e regiões."),
    ("", ""),
    ("CORES", ""),
    ("Texto azul / fundo amarelo", "Célula para você digitar ou escolher na lista."),
    ("Texto verde", "Valor puxado de outra aba (não digite)."),
    ("Texto preto", "Fórmula / cálculo automático (não digite)."),
    ("", ""),
    ("EXEMPLO", "A aba VENDAS traz um pedido de exemplo (Nº 1, marcado EXEMPLO). Apague as linhas 5 e 6 "
                "de VENDAS e a 1ª linha do registro de entregas (coluna H) em ENTREGADOR antes de começar a usar."),
    ("FONTE DOS PREÇOS", "Catálogo NOMAD puffs (docs/catalogo.pdf), informado pelos sócios em 01/10/2026."),
    ("GOOGLE SHEETS", "Arquivo > Importar > Fazer upload > Substituir planilha. Fórmulas e listas continuam funcionando."),
]
for i, (a, b) in enumerate(linhas, start=4):
    ws.cell(i, 1, a).font = Font(name=FONTE, size=10, bold=True,
                                 color=LARANJA if b == "" else "000000")
    ws.cell(i, 2, b).font = f_txt
    ws.cell(i, 2).alignment = Alignment(wrap_text=True, vertical="top")
ws["A13"].font, ws["B13"].fill = Font(name=FONTE, size=10, bold=True, color="0000FF"), fill_input
ws["A14"].font = Font(name=FONTE, size=10, bold=True, color="008000")

# ================================= VALOR ======================================
wv = wb.create_sheet("VALOR")
titulo(wv, "VALOR", "Preço de cada unidade por modelo", 5)
cabecalho(wv, 4, 1, ["Modelo", "Preço de venda (R$/un.)", "Custo (R$/un.)", "Lucro (R$/un.)", "Margem"],
          [22, 20, 16, 16, 12])
for i, (modelo, preco, _) in enumerate(MODELOS, start=5):
    cel(wv, f"A{i}", modelo)
    cel(wv, f"B{i}", preco, f_input, MOEDA, fill_input)
    cel(wv, f"C{i}", None, f_input, MOEDA, fill_input)
    cel(wv, f"D{i}", f'=IF(C{i}="","",B{i}-C{i})', fmt=MOEDA)
    cel(wv, f"E{i}", f'=IF(OR(C{i}="",B{i}=0),"",D{i}/B{i})', fmt="0.0%")
wv["B5"].comment = Comment("Preços do catálogo NOMAD puffs (docs/catalogo.pdf), 01/10/2026.", "NOMAD")
wv["C5"].comment = Comment("Preencha o custo de compra para calcular lucro e margem.", "NOMAD")
FIM_VALOR = 4 + len(MODELOS)
R_MODELOS = f"VALOR!$A$5:$A${FIM_VALOR}"
R_PRECOS = f"VALOR!$B$5:$B${FIM_VALOR}"
wv.freeze_panes = "A5"

# =============================== MERCADORIA ===================================
wm = wb.create_sheet("MERCADORIA")
titulo(wm, "MERCADORIA", "Estoque por modelo e sabor — vendido e saldo são automáticos", 9)
cabecalho(wm, 4, 1, ["Produto (código)", "Modelo", "Sabor", "Entradas (un.)", "Vendido (un.)",
                     "Estoque atual (un.)", "Valor unit. (R$)", "Valor em estoque (R$)", "Faturado (R$)"],
          [40, 18, 30, 14, 14, 16, 15, 19, 16])
r = 5
for modelo, _, sabores in MODELOS:
    for sabor in sabores:
        cel(wm, f"A{r}", f'=B{r}&" - "&C{r}')
        cel(wm, f"B{r}", modelo)
        cel(wm, f"C{r}", sabor)
        cel(wm, f"D{r}", None, f_input, INTEIRO, fill_input)
        cel(wm, f"E{r}", f"=SUMIFS({vcol("F")},{vcol("C")},A{r})", f_link, INTEIRO)
        cel(wm, f"F{r}", f"=N(D{r})-E{r}", fmt=INTEIRO)
        cel(wm, f"G{r}", f"=INDEX({R_PRECOS},MATCH(B{r},{R_MODELOS},0))", f_link, MOEDA)
        cel(wm, f"H{r}", f"=MAX(F{r},0)*G{r}", fmt=MOEDA)
        cel(wm, f"I{r}", f"=SUMIFS({vcol("H")},{vcol("C")},A{r})", f_link, MOEDA)
        r += 1
FIM_MERC = r - 1
for col, f in zip("ADEFHI", ["TOTAL", f"=SUM(D5:D{FIM_MERC})", f"=SUM(E5:E{FIM_MERC})",
                              f"=SUM(F5:F{FIM_MERC})", f"=SUM(H5:H{FIM_MERC})", f"=SUM(I5:I{FIM_MERC})"]):
    cel(wm, f"{col}{r}", f, f_total, MOEDA if col in "HI" else INTEIRO, fill_total)
for col in "BCG":
    cel(wm, f"{col}{r}", None, fill=fill_total)
wm.conditional_formatting.add(f"F5:F{FIM_MERC}", CellIsRule(
    operator="lessThan", formula=["0"], font=Font(color="C00000", bold=True),
    fill=PatternFill("solid", fgColor="F8CBAD")))
wm.conditional_formatting.add(f"F5:F{FIM_MERC}", FormulaRule(
    formula=[f'AND(F5>=0,F5<=3,N(D5)>0)'], fill=PatternFill("solid", fgColor="FFE699")))
wm.cell(r + 2, 1, "Vermelho = estoque negativo (vendeu mais do que entrou). "
                  "Amarelo = 3 unidades ou menos.").font = Font(name=FONTE, size=9, italic=True)
wm.freeze_panes = "B5"
R_PRODUTOS = f"MERCADORIA!$A$5:$A${FIM_MERC}"

# ================================== MAPA ======================================
wp = wb.create_sheet("MAPA")
titulo(wp, "MAPA — Entregas em Palmas – TO", "Taxa por quadra/bairro e desempenho de cada região", 8)
cabecalho(wp, 4, 1, ["Região de entrega", "Tipo", "Taxa de entrega (R$)", "Pedidos",
                     "Unidades", "Vendas em produtos (R$)", "Taxas recebidas (R$)", "% dos pedidos"],
          [34, 10, 18, 11, 11, 20, 18, 13])
for i, (regiao, tipo, taxa) in enumerate(REGIOES, start=5):
    cel(wp, f"A{i}", regiao)
    cel(wp, f"B{i}", tipo)
    cel(wp, f"C{i}", taxa, f_input, MOEDA, fill_input)
    cel(wp, f"D{i}", f"=SUMIFS({vcol("O")},{vcol("J")},A{i})", f_link, INTEIRO)
    cel(wp, f"E{i}", f"=SUMIFS({vcol("F")},{vcol("J")},A{i})", f_link, INTEIRO)
    cel(wp, f"F{i}", f"=SUMIFS({vcol("H")},{vcol("J")},A{i})", f_link, MOEDA)
    cel(wp, f"G{i}", f"=SUMIFS({vcol("K")},{vcol("J")},A{i})", f_link, MOEDA)
FIM_MAPA = 4 + len(REGIOES)
for i in range(5, FIM_MAPA + 1):
    cel(wp, f"H{i}", f"=IF($D${FIM_MAPA + 1}=0,0,D{i}/$D${FIM_MAPA + 1})", fmt="0.0%")
t = FIM_MAPA + 1
cel(wp, f"A{t}", "TOTAL", f_total, fill=fill_total)
for col, fmt in zip("DEFGH", [INTEIRO, INTEIRO, MOEDA, MOEDA, "0.0%"]):
    cel(wp, f"{col}{t}", f"=SUM({col}5:{col}{FIM_MAPA})", f_total, fmt, fill_total)
for col in "BC":
    cel(wp, f"{col}{t}", None, fill=fill_total)
wp["C5"].comment = Comment("Taxas do catálogo NOMAD puffs (docs/catalogo.pdf), 01/10/2026.", "NOMAD")
wp.cell(t + 2, 1, "Legenda das quadras (Plano Diretor de Palmas):").font = Font(name=FONTE, bold=True, size=10)
legenda = [
    "• 700 Sul a 200 Norte/Sul — região central (Praça dos Girassóis e entorno).",
    "• 300 a 600 Norte — Plano Diretor Norte.",
    "• 800 a 1200 — Plano Diretor Sul intermediário.",
    "• 1300 a 1500 Sul — extremo sul do Plano Diretor.",
    "• Bairros fora do Plano Diretor: Santo Amaro, Lago Norte, Bertaville, Aurenys, Taquaralto, Lago Sul, Taquari.",
    "• Região fora da lista: combinar a taxa no WhatsApp e lançar como 'Retirada / sem entrega' + ajustar a taxa.",
]
for k, txt in enumerate(legenda, start=t + 3):
    wp.cell(k, 1, txt).font = Font(name=FONTE, size=9)
grafico = BarChart()
grafico.type, grafico.style = "bar", 10
grafico.title, grafico.y_axis.title = "Pedidos por região", "Pedidos"
grafico.add_data(Reference(wp, min_col=4, min_row=4, max_row=FIM_MAPA), titles_from_data=True)
grafico.set_categories(Reference(wp, min_col=1, min_row=5, max_row=FIM_MAPA))
grafico.legend, grafico.height, grafico.width = None, 9, 16
wp.add_chart(grafico, "J4")
wp.freeze_panes = "A5"
R_REGIOES = f"MAPA!$A$5:$A${FIM_MAPA}"
R_TAXAS = f"MAPA!$C$5:$C${FIM_MAPA}"

# ============================== ENTREGADOR (cadastro) =========================
we = wb.create_sheet("ENTREGADOR")

# ================================= VENDAS =====================================
wvd = wb.create_sheet("VENDAS")
titulo(wvd, "VENDAS", "Uma linha por item vendido · totais somam sozinhos", 15)
cabecalho(wvd, 4, 1, ["Nº pedido", "Data", "Produto (modelo - sabor)", "Modelo", "Sabor", "Qtd",
                      "Valor unit. (R$)", "Subtotal (R$)", "Pagamento", "Região de entrega",
                      "Taxa entrega (R$)", "Total (R$)", "Entregador", "Endereço / observação",
                      "Pedido único"],
          [10, 12, 40, 16, 28, 7, 14, 14, 18, 32, 15, 14, 16, 34, 9])
for i in range(5, FIM_VENDAS + 1):
    for col in "ABCFIJMN":
        c = wvd[f"{col}{i}"]
        c.font, c.fill, c.border = f_input, fill_input, borda
    wvd[f"B{i}"].number_format = DATA
    wvd[f"D{i}"] = f'=IF(C{i}="","",INDEX(MERCADORIA!$B$5:$B${FIM_MERC},MATCH(C{i},MERCADORIA!$A$5:$A${FIM_MERC},0)))'
    wvd[f"E{i}"] = f'=IF(C{i}="","",INDEX(MERCADORIA!$C$5:$C${FIM_MERC},MATCH(C{i},MERCADORIA!$A$5:$A${FIM_MERC},0)))'
    wvd[f"G{i}"] = f'=IF(D{i}="","",INDEX({R_PRECOS},MATCH(D{i},{R_MODELOS},0)))'
    wvd[f"H{i}"] = f'=IF(OR(G{i}="",F{i}=""),0,F{i}*G{i})'
    # 1ª linha do pedido = conta como pedido e cobra a taxa uma única vez
    wvd[f"O{i}"] = f'=IF(A{i}="",0,IF(COUNTIF($A$5:A{i},A{i})=1,1,0))'
    wvd[f"K{i}"] = f'=IF(OR(O{i}=0,J{i}=""),0,INDEX({R_TAXAS},MATCH(J{i},{R_REGIOES},0)))'
    wvd[f"L{i}"] = f"=H{i}+K{i}"
    for col, fmt, fonte in [("D", None, f_link), ("E", None, f_link), ("G", MOEDA, f_link),
                            ("H", MOEDA, f_txt), ("K", MOEDA, f_link), ("L", MOEDA, f_total),
                            ("O", INTEIRO, f_txt)]:
        c = wvd[f"{col}{i}"]
        c.font, c.border = fonte, borda
        if fmt:
            c.number_format = fmt
    wvd[f"F{i}"].number_format = INTEIRO
lista(wvd, R_PRODUTOS, f"C5:C{FIM_VENDAS}", "Modelo - sabor")
lista(wvd, '"' + ",".join(PAGAMENTOS) + '"', f"I5:I{FIM_VENDAS}", "PIX, débito ou crédito")
lista(wvd, R_REGIOES, f"J5:J{FIM_VENDAS}", "Quadra/bairro da entrega")
lista(wvd, "ENTREGADOR!$A$5:$A$14", f"M5:M{FIM_VENDAS}", "Motoboy que levou o pedido")
qtd = DataValidation(type="whole", operator="between", formula1="1", formula2="100",
                     showErrorMessage=True, error="Quantidade entre 1 e 100.")
wvd.add_data_validation(qtd)
qtd.add(f"F5:F{FIM_VENDAS}")
wvd["O4"].comment = Comment("1 = primeira linha do pedido (conta o pedido e a taxa uma vez só).", "NOMAD")
# pedido de exemplo (2 itens)
exemplo = [
    (1, "2026-10-01", "V155 - Grape Ice", 2, "PIX", "Quadras 300 Norte a 600 Norte", "Motoboy 1",
     "EXEMPLO — apagar"),
    (1, "2026-10-01", "Elfbar Pro 40K - Pink Lemonade", 1, "PIX", "Quadras 300 Norte a 600 Norte",
     "Motoboy 1", "EXEMPLO — apagar"),
]
for i, (ped, dt, prod, q, pag, reg, ent, obs) in enumerate(exemplo, start=5):
    wvd[f"A{i}"], wvd[f"B{i}"] = ped, date.fromisoformat(dt)
    wvd[f"C{i}"], wvd[f"F{i}"], wvd[f"I{i}"] = prod, q, pag
    wvd[f"J{i}"], wvd[f"M{i}"], wvd[f"N{i}"] = reg, ent, obs
# linha de totais no topo (row 3) para ficar sempre visível
wvd["J3"] = "TOTAIS ▸"
wvd["F3"] = f"=SUM(F5:F{FIM_VENDAS})"
wvd["H3"] = f"=SUM(H5:H{FIM_VENDAS})"
wvd["K3"] = f"=SUM(K5:K{FIM_VENDAS})"
wvd["L3"] = f"=SUM(L5:L{FIM_VENDAS})"
wvd["O3"] = f"=SUM(O5:O{FIM_VENDAS})"
for col in "FHJKLO":
    c = wvd[f"{col}3"]
    c.font, c.fill, c.border = f_total, fill_total, borda
    c.number_format = MOEDA if col in "HKL" else INTEIRO
wvd.freeze_panes = "D5"
wvd.auto_filter.ref = f"A4:O{FIM_VENDAS}"

# ============================== ENTREGADOR ====================================
titulo(we, "ENTREGADOR", "Motoboys, entregas e modelos entregues", 20)
modelos_nomes = [m[0] for m in MODELOS]
FIM_ENT = 4 + LINHAS_ENTREGAS
# registro de entregas: colunas H..T (G é espaçador, A..F ficam para cadastro e resumo)
cab_reg = ["Nº pedido", "Data", "Motoboy", "Região", "Unidades"] + modelos_nomes +           ["Valor pedido (R$)", "Taxa entrega (R$)", "Status", "Horário entrega"]
RC = {nome: get_column_letter(8 + k) for k, nome in enumerate(cab_reg)}
PED, STA = RC["Nº pedido"], RC["Status"]


def reg(nome):
    """Intervalo de uma coluna do registro de entregas."""
    return f"${RC[nome]}$5:${RC[nome]}${FIM_ENT}"


# cadastro
cabecalho(we, 4, 1, ["Motoboy", "Telefone", "Pedidos entregues", "Unidades"], [18, 16, 12, 11])
for i in range(5, 15):
    nome = ENTREGADORES[i - 5] if i - 5 < len(ENTREGADORES) else None
    cel(we, f"A{i}", nome, f_input, fill=fill_input)
    cel(we, f"B{i}", None, f_input, fill=fill_input)
    cel(we, f"C{i}", f'=IF(A{i}="","",COUNTIFS({reg("Motoboy")},A{i},{reg("Status")},"Entregue"))',
        fmt=INTEIRO)
    cel(we, f"D{i}", f'=IF(A{i}="","",SUMIFS({reg("Unidades")},{reg("Motoboy")},A{i},'
                     f'{reg("Status")},"Entregue"))', fmt=INTEIRO)
we["A16"] = "Troque 'Motoboy 1/2/3' pelos nomes reais — as listas de VENDAS se atualizam."
we["A16"].font = Font(name=FONTE, size=9, italic=True)
for col, w in zip("EFG", [11, 13, 3]):
    we.column_dimensions[col].width = w
cabecalho(we, 4, 8, cab_reg, [10, 12, 14, 30, 10, 9, 9, 13, 14, 15, 15, 12, 13])
for i in range(5, FIM_ENT + 1):
    for nome in ("Nº pedido", "Status", "Horário entrega"):
        c = we[f"{RC[nome]}{i}"]
        c.font, c.fill, c.border = f_input, fill_input, borda
    p = f"{PED}{i}"
    m = f'MATCH({p},{vcol("A")},0)'
    formulas_reg = {
        "Data": (f'=IF({p}="","",IFERROR(INDEX({vcol("B")},{m}),"não lançado"))', DATA),
        "Motoboy": (f'=IF({p}="","",IFERROR(INDEX({vcol("M")},{m})&"",""))', None),
        "Região": (f'=IF({p}="","",IFERROR(INDEX({vcol("J")},{m})&"",""))', None),
        "Unidades": (f'=IF({p}="","",SUMIFS({vcol("F")},{vcol("A")},{p}))', INTEIRO),
        "Valor pedido (R$)": (f'=IF({p}="","",SUMIFS({vcol("H")},{vcol("A")},{p}))', MOEDA),
        "Taxa entrega (R$)": (f'=IF({p}="","",SUMIFS({vcol("K")},{vcol("A")},{p}))', MOEDA),
    }
    for modelo in modelos_nomes:
        formulas_reg[modelo] = (f'=IF({p}="","",SUMIFS({vcol("F")},{vcol("A")},{p},'
                                f'{vcol("D")},{RC[modelo]}$4))', INTEIRO)
    for nome, (f, fmt) in formulas_reg.items():
        c = we[f"{RC[nome]}{i}"]
        c.value, c.font, c.border = f, f_link, borda
        if fmt:
            c.number_format = fmt
    we[f"{RC['Horário entrega']}{i}"].number_format = "hh:mm"
lista(we, '"' + ",".join(STATUS) + '"', f"{STA}5:{STA}{FIM_ENT}", "Situação da entrega")
we[f"{PED}5"], we[f"{STA}5"] = 1, "Entregue"      # exemplo (pedido 1 de VENDAS)
# resumo motoboy x modelo (entregas com status Entregue)
lin = 18
we.cell(lin, 1, "UNIDADES ENTREGUES POR MODELO").font = Font(name=FONTE, size=10, bold=True, color=LARANJA)
cabecalho(we, lin + 1, 1, ["Motoboy"] + modelos_nomes + ["Taxas (R$)"])
for k in range(10):
    rr = lin + 2 + k
    cel(we, f"A{rr}", f'=IF(A{5 + k}="","",A{5 + k})', f_link)
    for col, modelo in zip("BCDE", modelos_nomes):
        cel(we, f"{col}{rr}", f'=IF($A{rr}="","",SUMIFS({reg(modelo)},{reg("Motoboy")},$A{rr},'
                              f'{reg("Status")},"Entregue"))', fmt=INTEIRO)
    cel(we, f"F{rr}", f'=IF($A{rr}="","",SUMIFS({reg("Taxa entrega (R$)")},{reg("Motoboy")},$A{rr},'
                      f'{reg("Status")},"Entregue"))', fmt=MOEDA)
we.cell(lin + 13, 1, "Taxas = soma das taxas de entrega dos pedidos com status 'Entregue' "
                     "(base para o repasse ao motoboy).").font = Font(name=FONTE, size=9, italic=True)
we.freeze_panes = "A5"

# ================================= RESUMO =====================================
wr = wb.create_sheet("RESUMO")
titulo(wr, "RESUMO", "Atualiza sozinho conforme VENDAS é preenchida", 6)
for col, w in zip("ABCDEF", [30, 18, 4, 22, 14, 18]):
    wr.column_dimensions[col].width = w
kpis = [
    ("Pedidos", f"=SUM(VENDAS!O5:O{FIM_VENDAS})", INTEIRO),
    ("Unidades vendidas", f"=SUM(VENDAS!F5:F{FIM_VENDAS})", INTEIRO),
    ("Vendas em produtos", f"=SUM(VENDAS!H5:H{FIM_VENDAS})", MOEDA),
    ("Taxas de entrega", f"=SUM(VENDAS!K5:K{FIM_VENDAS})", MOEDA),
    ("Faturamento total", "=B7+B8", MOEDA),
    ("Ticket médio por pedido", "=IF(B5=0,0,B9/B5)", MOEDA),
]
cabecalho(wr, 4, 1, ["Indicador", "Valor"])
for i, (nome, f, fmt) in enumerate(kpis, start=5):
    cel(wr, f"A{i}", nome, f_total)
    cel(wr, f"B{i}", f, f_total, fmt, fill_total)
# por pagamento
cabecalho(wr, 4, 4, ["Forma de pagamento", "Pedidos", "Valor (R$)"])
for i, pag in enumerate(PAGAMENTOS, start=5):
    cel(wr, f"D{i}", pag)
    cel(wr, f"E{i}", f"=SUMIFS(VENDAS!$O$5:$O${FIM_VENDAS},VENDAS!$I$5:$I${FIM_VENDAS},D{i})", fmt=INTEIRO)
    cel(wr, f"F{i}", f"=SUMIFS(VENDAS!$L$5:$L${FIM_VENDAS},VENDAS!$I$5:$I${FIM_VENDAS},D{i})", fmt=MOEDA)
cel(wr, "D8", "Cartão (débito + crédito)", f_total, fill=fill_total)
cel(wr, "E8", "=E6+E7", f_total, INTEIRO, fill_total)
cel(wr, "F8", "=F6+F7", f_total, MOEDA, fill_total)
cel(wr, "D9", "TOTAL", f_total, fill=fill_total)
cel(wr, "E9", "=E5+E8", f_total, INTEIRO, fill_total)
cel(wr, "F9", "=F5+F8", f_total, MOEDA, fill_total)
cel(wr, "D10", "Sem forma de pagamento", Font(name=FONTE, size=9, italic=True))
cel(wr, "E10", "=B5-E9", fmt=INTEIRO)
cel(wr, "F10", "=B9-F9", fmt=MOEDA)
# por modelo
cabecalho(wr, 13, 1, ["Modelo", "Unidades", "", "Vendas (R$)"])
for i, (modelo, _, _) in enumerate(MODELOS, start=14):
    cel(wr, f"A{i}", modelo)
    cel(wr, f"B{i}", f"=SUMIFS(VENDAS!$F$5:$F${FIM_VENDAS},VENDAS!$D$5:$D${FIM_VENDAS},A{i})", fmt=INTEIRO)
    cel(wr, f"D{i}", f"=SUMIFS(VENDAS!$H$5:$H${FIM_VENDAS},VENDAS!$D$5:$D${FIM_VENDAS},A{i})", fmt=MOEDA)
fm = 13 + len(MODELOS)
cel(wr, f"A{fm + 1}", "TOTAL", f_total, fill=fill_total)
cel(wr, f"B{fm + 1}", f"=SUM(B14:B{fm})", f_total, INTEIRO, fill_total)
cel(wr, f"D{fm + 1}", f"=SUM(D14:D{fm})", f_total, MOEDA, fill_total)
wr.cell(fm + 2, 1, "Conferência: TOTAL de modelos deve bater com 'Vendas em produtos'.").font = \
    Font(name=FONTE, size=9, italic=True)
wr.conditional_formatting.add(f"D{fm + 1}", FormulaRule(formula=[f"ROUND(D{fm + 1}-$B$7,2)<>0"],
                                                       fill=PatternFill("solid", fgColor="F8CBAD")))
pizza = PieChart()
pizza.title = "Pedidos por forma de pagamento"
pizza.add_data(Reference(wr, min_col=5, min_row=4, max_row=7), titles_from_data=True)
pizza.set_categories(Reference(wr, min_col=4, min_row=5, max_row=7))
pizza.height, pizza.width = 8, 12
wr.add_chart(pizza, "H4")
barras = BarChart()
barras.title, barras.style = "Unidades por modelo", 10
barras.add_data(Reference(wr, min_col=2, min_row=13, max_row=fm), titles_from_data=True)
barras.set_categories(Reference(wr, min_col=1, min_row=14, max_row=fm))
barras.legend, barras.height, barras.width = None, 8, 12
wr.add_chart(barras, "H21")

# ordem e cores das abas
wb._sheets = [wb["INÍCIO"], wb["RESUMO"], wb["VENDAS"], wb["MERCADORIA"], wb["VALOR"],
              wb["ENTREGADOR"], wb["MAPA"]]
for nome, cor in [("INÍCIO", PRETO), ("RESUMO", LARANJA), ("VENDAS", LARANJA),
                  ("MERCADORIA", CREME), ("VALOR", CREME), ("ENTREGADOR", CINZA), ("MAPA", CINZA)]:
    wb[nome].sheet_properties.tabColor = cor
wb.active = 0
wb.save(SAIDA)
print(f"OK -> {SAIDA}")
